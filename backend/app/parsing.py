import csv
import io
import re
from datetime import datetime, date
from zipfile import ZipFile
from openpyxl import load_workbook
from pypdf import PdfReader

ALIASES = {
    'item': 'name', 'item name': 'name', 'material': 'name', 'activity name': 'name',
    'lead time': 'lead_time_days', 'lead time days': 'lead_time_days', 'required on site': 'required_on_site_date',
    'required onsite': 'required_on_site_date', 'expected delivery': 'expected_delivery_date', 'release date': 'release_date',
    'start date': 'start_date', 'finish date': 'finish_date', 'activity id': 'activity_id',
    'percent complete': 'percent_complete', '% complete': 'percent_complete', 'submittal': 'submittal',
    'predecessors': 'predecessor_ids', 'successors': 'successor_ids',
}


def parse_table(content, filename, kind):
    if filename.lower().endswith('.csv'):
        text = content.decode('utf-8-sig')
        rows = list(csv.reader(io.StringIO(text)))
    elif filename.lower().endswith('.xlsx'):
        with ZipFile(io.BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist()) > 50_000_000:
                raise ValueError('Workbook expands beyond the 50 MB limit.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            rows = []
            for row in workbook.active.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > 10001:
                    raise ValueError('Use a workbook with at most 10,000 rows.')
        finally:
            workbook.close()
    else:
        raise ValueError('Please upload a CSV or XLSX file.')
    if len(rows) < 2 or len(rows) > 10001:
        raise ValueError('The file must contain headers and 1-10,000 data rows.')
    headers = [re.sub(r'\s+', ' ', str(h or '').replace('_', ' ').strip().lower()) for h in rows[0]]
    headers = [ALIASES.get(h, h.replace(' ', '_')) for h in headers]
    required = {'name', 'start_date', 'finish_date'} if kind == 'schedule' else {'name', 'status', 'required_on_site_date', 'lead_time_days'}
    if not required.issubset(headers):
        raise ValueError('Missing columns: ' + ', '.join(sorted(required - set(headers))))
    if len(headers) != len(set(headers)):
        raise ValueError('Duplicate column names are not supported.')
    result = []
    for line, cells in enumerate(rows[1:], 2):
        if not any(v is not None and str(v).strip() for v in cells):
            continue
        row = {h: v for h, v in zip(headers, cells) if h}
        try:
            if not str(row.get('name') or '').strip():
                raise ValueError('Item/activity name is required')
            for key, val in list(row.items()):
                if val is None or val == '':
                    row[key] = None
                elif key.endswith('_date'):
                    if isinstance(val, (datetime, date)):
                        row[key] = val.isoformat()[:10]
                    else:
                        parsed = None
                        for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y'):
                            try:
                                parsed = datetime.strptime(str(val).strip(), fmt).date().isoformat()
                                break
                            except ValueError:
                                pass
                        if not parsed:
                            raise ValueError(f'{key} must be YYYY-MM-DD or MM/DD/YYYY')
                        row[key] = parsed
                elif key == 'lead_time_days':
                    row[key] = int(str(val).strip())
                    if row[key] < 0 or row[key] > 3650:
                        raise ValueError('Lead time must be 0-3650 days')
                elif key == 'percent_complete':
                    row[key] = float(str(val).replace('%', ''))
                    if not 0 <= row[key] <= 100:
                        raise ValueError('Percent complete must be 0-100')
                elif key.endswith('_ids'):
                    row[key] = [x.strip() for x in str(val).split(',') if x.strip()]
                else:
                    row[key] = str(val).strip()
        except (ValueError, TypeError) as exc:
            raise ValueError(f'Row {line}: {exc}') from exc
        result.append(row)
    return result


def parse_document(content, filename):
    if filename.lower().endswith('.pdf'):
        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 300:
            raise ValueError('Please split documents larger than 300 pages.')
        pages = [(i + 1, p.extract_text() or '') for i, p in enumerate(reader.pages)]
    elif filename.lower().endswith('.txt'):
        pages = [(1, content.decode('utf-8-sig'))]
    else:
        raise ValueError('Please upload a PDF or UTF-8 TXT document.')
    if not any(text.strip() for _, text in pages):
        raise ValueError('No readable text was found. Run OCR on the scanned PDF and upload it again.')
    if sum(len(text) for _, text in pages) > 2_000_000:
        raise ValueError('Extracted text exceeds the 2 MB document limit.')
    return pages
