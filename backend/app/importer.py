"""Bring historical project data in from Excel or CSV exports.

The flow is two steps so nothing is written until the person has seen what will happen:
`analyze` reads every sheet, finds the header row, guesses what kind of records each sheet holds and
which column is which, and reports what would be imported. `commit` applies the same plan.
"""
import csv
import io
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import ZipFile
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from .database import get_db
from .models import Activity, Procurement, RFI, Submittal, Audit
from .risk_engine import run_rules
from .security import current_user, project_access, require_write

router = APIRouter(prefix='/api/projects/{project_id}/import')
Db, Auth = Depends(get_db), Depends(current_user)

MAX_ROWS = 20000
MAX_BYTES = 15_000_000

# field -> (label, kind of value, accepted header names in order of preference)
KINDS = {
    'schedule': {'label': 'Schedule activities', 'model': Activity, 'ident': ['name'], 'fields': {
        'name': ('Activity name', 'text', ['activity name', 'task name', 'activity', 'task', 'name', 'activity description', 'work item', 'description', 'item']),
        'start_date': ('Start', 'date', ['start date', 'planned start', 'early start', 'baseline start', 'start', 'begin']),
        'finish_date': ('Finish', 'date', ['finish date', 'planned finish', 'early finish', 'baseline finish', 'finish', 'end date', 'end', 'completion date', 'due date']),
        'activity_id': ('Activity ID', 'text', ['activity id', 'task id', 'activity code', 'wbs', 'id', 'code']),
        'trade': ('Trade', 'text', ['trade', 'responsible party', 'responsible', 'subcontractor', 'contractor', 'crew', 'discipline', 'assigned to']),
        'percent_complete': ('% complete', 'percent', ['percent complete', '% complete', 'pct complete', 'progress', 'complete', 'percent']),
        'predecessor_ids': ('Predecessors', 'list', ['predecessors', 'predecessor']),
        'successor_ids': ('Successors', 'list', ['successors', 'successor']),
    }},
    'procurement': {'label': 'Procurement / deliveries', 'model': Procurement, 'ident': ['name'], 'fields': {
        'name': ('Item', 'text', ['item name', 'material', 'equipment', 'product', 'procurement item', 'line item', 'item', 'name', 'description']),
        'vendor': ('Vendor', 'text', ['vendor', 'supplier', 'manufacturer', 'company']),
        'status': ('Status', 'status', ['order status', 'delivery status', 'status', 'state']),
        'lead_time_days': ('Lead time', 'days', ['lead time days', 'lead time', 'lead']),
        'required_on_site_date': ('Required on site', 'date', ['required on site date', 'required on site', 'required onsite', 'need by', 'needed by', 'need date', 'required date', 'date required', 'on site date', 'rosd']),
        'expected_delivery_date': ('Expected delivery', 'date', ['expected delivery date', 'expected delivery', 'estimated delivery', 'delivery date', 'promised date', 'eta', 'ship date', 'delivery']),
        'release_date': ('Release / order date', 'date', ['release date', 'order date', 'date ordered', 'po date', 'released']),
        'notes': ('Notes', 'text', ['notes', 'comments', 'remarks']),
    }},
    'rfis': {'label': 'RFIs', 'model': RFI, 'ident': ['subject', 'number'], 'fields': {
        'number': ('RFI number', 'text', ['rfi number', 'rfi #', 'rfi no', 'rfi', 'number', 'no', '#', 'id']),
        'subject': ('Subject', 'text', ['rfi subject', 'subject', 'title', 'name']),
        'question': ('Question', 'text', ['rfi question', 'question', 'details', 'description']),
        'response': ('Response', 'text', ['official response', 'rfi response', 'response', 'answer']),
        'status': ('Status', 'status', ['rfi status', 'status', 'state']),
        'due_date': ('Due date', 'date', ['response due', 'date due', 'due date', 'due', 'required response date']),
        'ball_in_court': ('Ball in court', 'text', ['ball in court', 'assigned to', 'assignee', 'responsible']),
    }},
    'submittals': {'label': 'Submittals', 'model': Submittal, 'ident': ['title', 'number'], 'fields': {
        'number': ('Submittal number', 'text', ['submittal number', 'submittal #', 'submittal no', 'number', 'no', '#', 'id']),
        'title': ('Title', 'text', ['submittal title', 'title', 'submittal', 'name', 'subject']),
        'description': ('Description', 'text', ['description', 'details']),
        'status': ('Status', 'status', ['submittal status', 'status', 'state']),
        'due_date': ('Review due', 'date', ['review due', 'date due', 'due date', 'due']),
        'required_on_site_date': ('Required on site', 'date', ['required on site date', 'required on site', 'required onsite', 'need by', 'needed by', 'required date']),
        'lead_time_days': ('Lead time', 'days', ['lead time days', 'lead time', 'lead']),
        'ball_in_court': ('Ball in court', 'text', ['ball in court', 'assigned to', 'assignee', 'responsible']),
        'spec_section': ('Spec section', 'text', ['spec section', 'specification section', 'spec', 'section', 'division']),
    }},
}
SHEET_HINTS = {'schedule': ('schedule', 'lookahead', 'look ahead', 'activit', 'task', 'gantt'), 'procurement': ('procure', 'deliver', 'material', 'equipment', 'purchase'), 'rfis': ('rfi',), 'submittals': ('submittal',)}
COMPLETE = {'schedule': {'percent_complete': 100.0}, 'procurement': {'status': 'delivered'}, 'rfis': {'status': 'closed'}, 'submittals': {'status': 'approved'}}
DATE_FORMATS = ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y', '%m-%d-%Y', '%d-%b-%Y', '%d-%b-%y', '%b %d, %Y', '%B %d, %Y', '%b %d %Y', '%d %b %Y', '%d %B %Y', '%Y/%m/%d', '%m.%d.%Y')
STATUS_WORDS = {'complete': 'delivered', 'completed': 'delivered', 'received': 'delivered', 'delivered': 'delivered', 'on site': 'delivered', 'onsite': 'delivered',
                'ordered': 'ordered', 'po issued': 'ordered', 'released': 'released', 'in transit': 'ordered', 'shipped': 'ordered'}


def norm(value):
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9%#]+', ' ', str(value if value is not None else '').lower())).strip()


def read_sheets(content, filename):
    name = filename.lower()
    if name.endswith('.csv') or name.endswith('.txt'):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('latin-1')
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        return [('Sheet1', list(csv.reader(io.StringIO(text), dialect)))]
    if name.endswith('.xls'):
        raise ValueError('That is an older .xls workbook. Open it in Excel, choose Save As, and pick Excel Workbook (.xlsx), then upload it again.')
    if not name.endswith('.xlsx'):
        raise ValueError('Please upload an Excel (.xlsx) or CSV file.')
    try:
        with ZipFile(io.BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist()) > 80_000_000:
                raise ValueError('That workbook is too large once opened. Split it into smaller files.')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ValueError:
        raise
    except Exception:
        raise ValueError('That file could not be opened as an Excel workbook. If it is password-protected, remove the password and try again.')
    sheets = []
    try:
        for ws in book.worksheets:
            if getattr(ws, 'sheet_state', 'visible') != 'visible':
                continue
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_ROWS + 50:
                    raise ValueError(f'Sheet "{ws.title}" has more than {MAX_ROWS:,} rows. Split it into smaller files.')
            sheets.append((ws.title, rows))
    finally:
        book.close()
    return sheets


def known_headers():
    return {name for k in KINDS.values() for f in k['fields'].values() for name in f[2]}


def find_header(rows):
    known = known_headers()
    best, best_score = None, 1
    for i, row in enumerate(rows[:25]):
        score = sum(1 for c in row if norm(c) in known)
        if score > best_score:
            best, best_score = i, score
    if best is not None:
        return best
    for i, row in enumerate(rows[:25]):
        if sum(1 for c in row if c not in (None, '')) >= 2:
            return i
    return None


def suggest_mapping(headers, kind):
    """field -> column index. Exact header names first, then headers that contain a known name."""
    names = [norm(h) for h in headers]
    taken, mapping = set(), {}
    fields = KINDS[kind]['fields']
    for contains in (False, True):
        for field, (_, _, synonyms) in fields.items():
            if field in mapping:
                continue
            for syn in synonyms:
                if contains and len(syn) < 4:
                    continue
                for i, n in enumerate(names):
                    if i in taken or not n:
                        continue
                    if n == syn or (contains and re.search(r'(^| )' + re.escape(syn) + r'( |$)', n)):
                        mapping[field] = i
                        taken.add(i)
                        break
                if field in mapping:
                    break
    return mapping


def detect_kind(headers, sheet_name):
    best, best_score = None, 0
    lowered = sheet_name.lower()
    for kind, spec in KINDS.items():
        mapping = suggest_mapping(headers, kind)
        if not any(f in mapping for f in spec['ident']):
            continue
        score = len(mapping) + (2 if any(h in lowered for h in SHEET_HINTS[kind]) else 0)
        if score > best_score:
            best, best_score = kind, score
    return best if best_score >= 3 else None


def to_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 < value < 80000:  # Excel serial day number
            return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
        raise ValueError('not a date')
    text = str(value).strip()
    text = re.sub(r'^(mon|tue|wed|thu|fri|sat|sun)[a-z]*,?\s+', '', text, flags=re.I)
    text = re.sub(r'(\d)(st|nd|rd|th)\b', r'\1', text)
    text = re.split(r'[T ](?=\d{1,2}:\d{2})', text)[0]
    for fmt in DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt).date()
            if 1990 <= parsed.year <= 2100:
                return parsed.isoformat()
        except ValueError:
            pass
    raise ValueError('not a date')


def to_days(value):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
    else:
        m = re.match(r'^\s*(\d+(?:\.\d+)?)\s*([a-z]*)', str(value).lower())
        if not m:
            raise ValueError('not a number of days')
        number = float(m.group(1))
        unit = m.group(2)
        if unit.startswith('w'):
            number *= 7
        elif unit.startswith('m') and unit not in ('min', 'mins'):
            number *= 30
    if not 0 <= number <= 3650:
        raise ValueError('out of range')
    return int(round(number))


def to_status(value):
    text = norm(value)
    return STATUS_WORDS.get(text, text) or None


def clean_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() or None


def convert(rows, header_row, kind, mapping, assume_complete=False):
    spec = KINDS[kind]
    fields = spec['fields']
    mapping = {f: int(i) for f, i in (mapping or {}).items() if f in fields and i is not None and str(i) != ''}
    percent_scale = 1
    if 'percent_complete' in mapping:
        col = [r[mapping['percent_complete']] for r in rows[header_row + 1:] if mapping['percent_complete'] < len(r)]
        nums = [v for v in col if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if nums and max(nums) <= 1 and any(isinstance(v, float) and 0 < v < 1 for v in nums):
            percent_scale = 100
    records, issues, skipped = [], [], 0
    for line, cells in enumerate(rows[header_row + 1:], header_row + 2):
        if not any(c is not None and str(c).strip() for c in cells):
            continue
        record, problems = {}, []
        for field, index in mapping.items():
            raw = cells[index] if index < len(cells) else None
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                continue
            label, typ = fields[field][0], fields[field][1]
            try:
                if typ == 'date':
                    record[field] = to_date(raw)
                elif typ == 'days':
                    record[field] = to_days(raw)
                elif typ == 'percent':
                    number = float(str(raw).replace('%', '').strip()) * (percent_scale if isinstance(raw, (int, float)) else 1)
                    record[field] = max(0.0, min(100.0, number))
                elif typ == 'list':
                    record[field] = [x.strip() for x in re.split(r'[,;]', str(raw)) if x.strip()]
                elif typ == 'status':
                    record[field] = to_status(raw)
                else:
                    record[field] = clean_text(raw)
            except (ValueError, TypeError):
                problems.append(f'{label} "{str(raw)[:30]}" could not be read, left blank')
        if not any(record.get(f) for f in spec['ident']):
            skipped += 1
            if len(issues) < 25:
                issues.append(f'Row {line}: skipped, no {" or ".join(fields[f][0].lower() for f in spec["ident"])}')
            continue
        if kind == 'schedule' and not record.get('name'):
            continue
        if kind == 'rfis' and not record.get('subject'):
            record['subject'] = f'RFI {record["number"]}'
        if kind == 'submittals' and not record.get('title'):
            record['title'] = f'Submittal {record["number"]}'
        if assume_complete:
            for key, value in COMPLETE[kind].items():
                record.setdefault(key, value)
        for p in problems:
            if len(issues) < 25:
                issues.append(f'Row {line}: {p}')
        records.append(record)
    return records, issues, skipped


def natural_key(kind, r):
    if kind == 'schedule':
        return 'id:' + norm(r['activity_id']) if r.get('activity_id') else 'n:' + norm(r['name']) + '|' + (r.get('start_date') or '')
    if kind == 'procurement':
        return norm(r['name']) + '|' + norm(r.get('vendor')) + '|' + (r.get('required_on_site_date') or '')
    return norm(r.get('number')) or norm(r.get('subject') or r.get('title'))


def analyze(content, filename, plan=None, assume_complete=False):
    """plan: optional {sheet name: {'kind': str | None, 'mapping': {field: column index}}} overrides the guesses."""
    plan = plan or {}
    out = []
    for name, rows in read_sheets(content, filename):
        header_row = find_header(rows)
        if header_row is None:
            continue
        headers = ['' if h is None else str(h).strip() for h in rows[header_row]]
        mine = plan.get(name) or {}
        kind = mine['kind'] if 'kind' in mine else detect_kind(headers, name)
        if kind not in KINDS:
            kind = None
        mapping = mine.get('mapping') if kind and mine.get('mapping') is not None and mine.get('kind') == kind else (suggest_mapping(headers, kind) if kind else {})
        entry = {'sheet': name, 'header_row': header_row + 1, 'headers': headers, 'kind': kind, 'mapping': mapping,
                 'row_count': sum(1 for r in rows[header_row + 1:] if any(c is not None and str(c).strip() for c in r))}
        if kind:
            records, issues, skipped = convert(rows, header_row, kind, mapping, assume_complete)
            unique = {natural_key(kind, r): r for r in records}
            entry.update(importable=len(unique), duplicates=len(records) - len(unique), skipped=skipped, issues=issues, sample=records[:5])
            entry['records'] = list(unique.values())
        out.append(entry)
    if not out:
        raise ValueError('No rows were found in that file.')
    return out


def apply(db, project_id, kind, records, filename, sheet):
    spec = KINDS[kind]
    cls = spec['model']
    allowed = set(cls.__table__.columns.keys()) - {'id', 'project_id', 'created_at', 'source_external_id', 'raw_json'}
    existing = {r.source_external_id: r for r in db.scalars(select(cls).where(cls.project_id == project_id, cls.source_external_id.like('xl:%')))}
    created = updated = 0
    for record in records:
        key = 'xl:' + kind + ':' + natural_key(kind, record)
        row = existing.get(key)
        if row is None:
            row = cls(project_id=project_id, source_external_id=key)
            db.add(row)
            created += 1
        else:
            updated += 1
        for field, value in record.items():
            if field in allowed:
                setattr(row, field, value)
        row.raw_json = {**(row.raw_json or {}), **record, 'imported_from': filename, 'sheet': sheet}
    return created, updated


@router.post('/preview')
async def preview(project_id: str, file: UploadFile = File(...), plan: str = Form('{}'), assume_complete: bool = Form(False), db: Session = Db, user=Auth):
    require_write(user)
    project_access(db, user, project_id)
    content = await file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Please upload a file smaller than 15 MB.')
    try:
        sheets = analyze(content, Path(file.filename or 'upload').name, json.loads(plan or '{}'), assume_complete)
    except json.JSONDecodeError:
        raise HTTPException(422, 'The import plan could not be read.')
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    for s in sheets:
        s.pop('records', None)
    return {'filename': file.filename, 'sheets': sheets, 'kinds': {k: {'label': v['label'], 'ident': v['ident'], 'fields': {f: {'label': d[0], 'type': d[1]} for f, d in v['fields'].items()}} for k, v in KINDS.items()}}


@router.post('')
async def commit(project_id: str, file: UploadFile = File(...), plan: str = Form('{}'), assume_complete: bool = Form(False), db: Session = Db, user=Auth):
    require_write(user)
    project_access(db, user, project_id)
    content = await file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Please upload a file smaller than 15 MB.')
    filename = Path(file.filename or 'upload').name[:200]
    try:
        sheets = analyze(content, filename, json.loads(plan or '{}'), assume_complete)
    except json.JSONDecodeError:
        raise HTTPException(422, 'The import plan could not be read.')
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    summary = []
    for s in sheets:
        if not s['kind'] or not s.get('records'):
            continue
        created, updated = apply(db, project_id, s['kind'], s['records'], filename, s['sheet'])
        summary.append({'sheet': s['sheet'], 'kind': s['kind'], 'created': created, 'updated': updated, 'skipped': s['skipped']})
    if not summary:
        raise HTTPException(422, 'Nothing to import. Pick what each sheet contains and check the columns.')
    db.flush()
    run_rules(db, project_id)
    db.add(Audit(organization_id=user.organization_id, user_id=user.id, action='import.history', detail={'project_id': project_id, 'filename': filename, 'sheets': summary}))
    db.commit()
    total = sum(x['created'] + x['updated'] for x in summary)
    return {'imported': total, 'created': sum(x['created'] for x in summary), 'updated': sum(x['updated'] for x in summary), 'sheets': summary, 'message': f'Imported {total} records from {filename}.'}
